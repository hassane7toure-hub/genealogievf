"""Génère le guide de procédure utilisateur au format Word."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "Procedure_utilisateur_TOURE_FAMILY_HERITAGE.docx"

UMBER = RGBColor(0x5C, 0x3A, 0x21)
MUTED = RGBColor(0x5A, 0x4A, 0x3A)
BLACK = RGBColor(0x1F, 0x1A, 0x14)


def set_run(run, *, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, italic=False, color=BLACK, space_after=8, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading_custom(doc, text, level):
    sizes = {1: 20, 2: 15, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run(run, size=sizes[level], bold=True, color=UMBER)
    return p


def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(6)
        num = p.add_run(f"{i}. ")
        set_run(num, size=11, bold=True, color=UMBER)
        body = p.add_run(step)
        set_run(body, size=11)


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): "5C3A21", qn("w:val"): "clear"})
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_header(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        set_run(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run(run, size=10)
    doc.add_paragraph()
    return table


def add_note(doc, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.3)
    t = p.add_run(f"{title} ")
    set_run(t, size=11, bold=True, color=UMBER)
    b = p.add_run(text)
    set_run(b, size=11, italic=True, color=MUTED)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("TOURÉ FAMILY HERITAGE — Modules de formation utilisateur et administrateur — Document interne")
    set_run(fr, size=8, color=MUTED)

    add_para(doc, "TOURÉ FAMILY HERITAGE", size=12, bold=True, color=UMBER, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(
        doc,
        "Modules de formation — utilisateur et administrateur",
        size=22,
        bold=True,
        color=UMBER,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_para(
        doc,
        "Module utilisateur : consulter, saisir une lignée, déposer un héritage. Module administrateur : promouvoir les profils, valider, publier.",
        size=12,
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=16,
    )
    add_para(doc, f"Version 2.0  ·  {date.today().strftime('%d/%m/%Y')}  ·  https://genealogievf.vercel.app  ·  formation locale http://localhost:3000", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    add_heading_custom(doc, "1. Objet du document", 1)
    add_para(
        doc,
        "Ce document comporte deux modules. Le module utilisateur explique comment un visiteur, un invité, un membre ou un descendant consulte l’arbre à partir d’une personne, ajoute une fiche, lie un conjoint et dépose une source, un témoignage ou un document. Le module administrateur explique comment le comité et l’administration promeuvent un profil, tranchent une validation avec un procès-verbal, publient une pièce publique et tiennent la caisse. Les libellés et les adresses correspondent à l’application actuelle.",
    )
    add_note(
        doc,
        "Principe de vérité :",
        "une information n’est pas un fait tant qu’elle n’est pas sourcée. On n’invente aucun lien de parenté. Les homonymes (deux Bourahima TOURE, Lanfia et Lankafia, Mâ Sokona et Mâ Sokonassé) restent des fiches distinctes jusqu’à validation du comité.",
    )

    add_heading_custom(doc, "2. Adresses utiles", 1)
    add_table(
        doc,
        ["Page", "Adresse", "Accès"],
        [
            ["Accueil public", "http://localhost:3000/", "Tout le monde, sans compte"],
            ["Histoire", "http://localhost:3000/histoire", "Public (C0)"],
            ["Généalogie publique", "http://localhost:3000/genealogie", "Public (C0)"],
            ["Arbre public", "http://localhost:3000/arbre", "Public (C0)"],
            ["Recherche publique", "http://localhost:3000/recherche", "Public (C0)"],
            ["Inscription", "http://localhost:3000/sign-up", "Nouveau compte"],
            ["Connexion", "http://localhost:3000/sign-in", "Compte existant"],
            ["Espace famille", "http://localhost:3000/espace", "Compte connecté"],
            ["Personnes (interne)", "http://localhost:3000/espace/personnes", "Compte connecté"],
            ["Nouvelle personne", "http://localhost:3000/espace/personnes/nouvelle", "Rôle d’écriture"],
            ["Arbre interne", "http://localhost:3000/espace/arbre", "Compte connecté"],
            ["Recherche interne", "http://localhost:3000/espace/recherche", "Compte connecté"],
            ["Réunions", "http://localhost:3000/espace/reunions", "Membre reconnu (C1)"],
            ["Événements", "http://localhost:3000/espace/evenements", "Membre reconnu (C1)"],
            ["Discussions", "http://localhost:3000/espace/discussions", "Membre reconnu (C1)"],
            ["Caisse", "http://localhost:3000/espace/caisse", "Comité et administration"],
            ["Notifications", "http://localhost:3000/espace/notifications", "Compte connecté"],
            ["Héritage — sources", "http://localhost:3000/espace/heritage", "Membre reconnu (C1)"],
            ["Témoignages", "http://localhost:3000/espace/heritage/temoignages", "Membre reconnu (C1)"],
            ["Documents", "http://localhost:3000/espace/heritage/documents", "Membre reconnu (C1)"],
            ["Validation", "http://localhost:3000/espace/heritage/validation", "Lecture C1 ; décision : comité"],
        ],
    )
    add_para(
        doc,
        "Les mêmes chemins fonctionnent sur le site publié https://genealogievf.vercel.app. Les pages /espace exigent une connexion : un visiteur non connecté est renvoyé vers /sign-in.",
    )

    add_heading_custom(doc, "3. Les types de profils", 1)
    add_para(
        doc,
        "L’application sépare trois notions. Ne pas les confondre : un descendant direct n’est pas automatiquement membre du comité, et un compte nouvellement créé n’écrit pas encore dans l’arbre.",
    )
    add_heading_custom(doc, "3.1 Rôle applicatif (ce que la personne PEUT faire)", 2)
    add_table(
        doc,
        ["Profil", "Rôle technique", "Ce qu’il voit", "Ce qu’il peut écrire"],
        [
            [
                "Visiteur (sans compte)",
                "—",
                "Uniquement les fiches publiques C0 (Lanfia, Samory, liste SRC-001, ascendance explicite).",
                "Rien. Consultation seulement.",
            ],
            [
                "Invité",
                "GUEST",
                "Comme le visiteur (C0), même après connexion, tant que le rôle n’a pas été élevé.",
                "Rien. Le bouton « Ajouter une personne » n’apparaît pas.",
            ],
            [
                "Membre de la Grande Famille",
                "FAMILY_MEMBER",
                "C0 et C1 (listes Sanankoro SRC-002 et SRC-003, vie familiale interne).",
                "Créer une fiche, lier un conjoint, animer la vie de famille, déposer un brouillon d’héritage et le soumettre. Pas la caisse, pas la décision du comité.",
            ],
            [
                "Descendant direct",
                "GUEST, FAMILY_MEMBER ou autre + statut DIRECT_DESCENDANT",
                "Jusqu’à C2 si le statut généalogique est « Descendant direct ».",
                "Selon le rôle. Le statut ouvre la lecture C2 ; il n’accorde pas à lui seul l’écriture.",
            ],
            [
                "Comité restreint",
                "COMMITTEE_MEMBER",
                "Jusqu’à C3 (informations sensibles).",
                "Tout ce que fait un membre, plus le procès-verbal, l’approbation, la publication C0 et la caisse.",
            ],
            [
                "Administration",
                "ADMIN",
                "Jusqu’à C3.",
                "Comme le comité. En plus : promouvoir les rôles (Prisma Studio ou variable d’amorçage).",
            ],
        ],
    )
    add_heading_custom(doc, "3.2 Statut généalogique (qui la personne EST dans la lignée)", 2)
    add_table(
        doc,
        ["Statut", "Signification"],
        [
            ["Non rattaché (UNLINKED)", "Compte créé, pas encore relié à une fiche Personne."],
            ["Membre de la lignée (FAMILY_MEMBER)", "Reconnu dans la famille, sans filiation directe validée."],
            ["Descendant direct (DIRECT_DESCENDANT)", "Filiation validée. Ouvre la lecture C2. La présence aux réunions n’est pas exigée."],
        ],
    )
    add_heading_custom(doc, "3.3 Confidentialité des fiches", 2)
    add_table(
        doc,
        ["Niveau", "Nature", "Qui peut lire"],
        [
            ["C0", "Informations publiques validées ou historiques", "Tout visiteur, sans e-mail"],
            ["C1", "Informations familiales internes", "Membre de la Grande Famille et au-delà"],
            ["C2", "Informations réservées aux descendants directs", "Statut descendant direct, comité, admin"],
            ["C3", "Informations sensibles", "Comité et administration seulement"],
        ],
    )

    add_heading_custom(doc, "4. Identifiants d’exemple par type de profil", 1)
    add_para(
        doc,
        "L’application n’embarque aucun identifiant d’usine et aucun mot de passe par défaut. L’authentification est déléguée à Clerk : chaque personne crée son e-mail et choisit son mot de passe à l’inscription. Le tableau ci-dessous est le jeu d’exemples à créer pour tester chaque profil. Ce ne sont pas des comptes déjà présents dans la base.",
    )
    add_table(
        doc,
        ["Profil à tester", "Identifiant (e-mail d’exemple)", "Mot de passe d’exemple", "Rôle à attribuer ensuite", "Statut généalogique"],
        [
            ["Visiteur", "Aucun", "Aucun", "—", "—"],
            [
                "Invité (compte fraîchement créé)",
                "invite.exemple@toure-heritage.test",
                "Heritage2026!Invite",
                "GUEST (valeur automatique à l’inscription)",
                "UNLINKED (automatique)",
            ],
            [
                "Membre de la Grande Famille",
                "famille.exemple@toure-heritage.test",
                "Heritage2026!Famille",
                "FAMILY_MEMBER",
                "FAMILY_MEMBER",
            ],
            [
                "Descendant direct",
                "descendant.exemple@toure-heritage.test",
                "Heritage2026!Lignee",
                "FAMILY_MEMBER",
                "DIRECT_DESCENDANT",
            ],
            [
                "Comité restreint",
                "comite.exemple@toure-heritage.test",
                "Heritage2026!Comite",
                "COMMITTEE_MEMBER",
                "FAMILY_MEMBER",
            ],
            [
                "Administration",
                "admin.exemple@toure-heritage.test",
                "Heritage2026!Admin",
                "ADMIN",
                "UNLINKED ou FAMILY_MEMBER",
            ],
        ],
    )
    add_note(
        doc,
        "Sécurité :",
        "changer ces mots de passe d’exemple dès qu’ils ont servi à une démonstration. Ne jamais les utiliser en production. Ne jamais commiter de vrais secrets dans Git.",
    )
    add_heading_custom(doc, "4.1 Comment un compte devient réellement ce profil", 2)
    add_para(doc, "À l’inscription, tout le monde devient Invité (GUEST) et Non rattaché (UNLINKED). Pour obtenir un autre profil :")
    add_steps(
        doc,
        [
            "La personne s’inscrit et se connecte une première fois (cela crée la ligne User dans PostgreSQL).",
            "Un administrateur ouvre Prisma Studio : dans un terminal, à la racine du projet, lancer npm run db:studio, puis ouvrir http://localhost:5555.",
            "Ouvrir le modèle User, trouver la ligne par e-mail.",
            "Modifier le champ role (GUEST, FAMILY_MEMBER, COMMITTEE_MEMBER ou ADMIN) et, si besoin, genealogicalStatus (UNLINKED, FAMILY_MEMBER ou DIRECT_DESCENDANT).",
            "Enregistrer. La personne se déconnecte puis se reconnecte, ou recharge /espace : le tableau de bord affiche le nouveau rôle et l’accès maximal (C0, C1, C2 ou C3).",
        ],
    )
    add_para(
        doc,
        "Le premier administrateur peut aussi être promu automatiquement si sa clé Clerk (user_…) est placée dans la variable d’environnement BOOTSTRAP_ADMIN_CLERK_USER_ID. Tant que son rôle est encore GUEST, la prochaine connexion le passe en ADMIN.",
    )

    add_heading_custom(doc, "5. Procédure d’inscription", 1)
    add_para(doc, "Valable pour tous les profils. Le rôle réel (famille, comité, admin) s’obtient après l’étape 4.1.")
    add_steps(
        doc,
        [
            "Ouvrir le navigateur et aller sur http://localhost:3000/.",
            "En haut à droite, cliquer sur le bouton « Rejoindre » (ou ouvrir directement http://localhost:3000/sign-up).",
            "Saisir l’adresse e-mail d’exemple du profil à créer (voir le tableau du § 4).",
            "Choisir le mot de passe d’exemple correspondant. Clerk exige en général 8 caractères au moins, avec une complexité minimale.",
            "Renseigner le prénom et le nom s’ils sont demandés (ex. « Awa » et « TOURÉ »).",
            "Valider l’inscription. Si Clerk envoie un code de vérification, l’ouvrir dans la boîte mail et le saisir.",
            "Après validation, l’application redirige vers /espace (Espace famille).",
            "Vérifier le bandeau gauche : le rôle affiché est « Invité » et le statut « Non rattaché », avec le badge « Accès C0 », jusqu’à promotion.",
        ],
    )
    add_note(
        doc,
        "Exemple concret — inscription d’un membre de la Grande Famille :",
        "e-mail famille.exemple@toure-heritage.test, mot de passe Heritage2026!Famille. Après la première connexion, l’administrateur passe le rôle à FAMILY_MEMBER. Au prochain chargement de /espace, le tableau de bord affiche « Membre de la Grande Famille » et « Accès C1 ». Les 139 fiches Sanankoro deviennent alors visibles.",
    )

    add_heading_custom(doc, "6. Procédure de connexion", 1)
    add_heading_custom(doc, "6.1 Visiteur (sans identifiant)", 2)
    add_steps(
        doc,
        [
            "Ouvrir http://localhost:3000/ sans cliquer sur Connexion.",
            "Naviguer dans Histoire, Généalogie, Arbre, Recherche.",
            "Constater que seules les fiches C0 apparaissent (environ 30 personnes après l’import Excel).",
            "Sur /arbre, rien ne s’affiche tant qu’une personne n’est pas choisie. Dans « Afficher l’arbre à partir de », filtrer un nom, choisir la personne, puis cliquer « Afficher ». L’arbre commence à cette personne et montre sa descendance. « Lignée complète » montre toute la lignée publique, de la racine jusqu’aux descendants visibles.",
        ],
    )
    add_heading_custom(doc, "6.2 Compte existant (tous les profils authentifiés)", 2)
    add_steps(
        doc,
        [
            "Ouvrir http://localhost:3000/.",
            "Cliquer sur « Connexion » en haut à droite (page /sign-in).",
            "Saisir l’identifiant et le mot de passe du profil (tableau § 4).",
            "Valider. La redirection mène à /espace.",
            "Contrôler le tableau de bord : carte « Autorisation » (rôle), statut généalogique, accès maximal, nombre de fiches visibles.",
            "Utiliser le menu : Tableau de bord, Personnes, Arbre, Réunions, Événements, Discussions, Caisse, Notifications, Héritage, Recherche, Site public. La caisse et les décisions de validation restent réservées au comité et à l’administration.",
            "Pour se déconnecter : cliquer sur l’avatar Clerk (UserButton) puis Sign out.",
        ],
    )
    add_heading_custom(doc, "6.3 Exemple de connexion par profil", 2)
    add_table(
        doc,
        ["Profil", "Saisir", "Résultat attendu sur /espace"],
        [
            [
                "Invité",
                "invite.exemple@toure-heritage.test / Heritage2026!Invite",
                "Rôle Invité, Accès C0, pas de bouton « Ajouter une personne ».",
            ],
            [
                "Membre famille",
                "famille.exemple@toure-heritage.test / Heritage2026!Famille",
                "Rôle Membre de la Grande Famille, Accès C1, bouton « Ajouter une personne » visible.",
            ],
            [
                "Descendant direct",
                "descendant.exemple@toure-heritage.test / Heritage2026!Lignee",
                "Statut Descendant direct, Accès C2, fiches C2 lisibles.",
            ],
            [
                "Comité",
                "comite.exemple@toure-heritage.test / Heritage2026!Comite",
                "Rôle Comité restreint, Accès C3.",
            ],
            [
                "Administration",
                "admin.exemple@toure-heritage.test / Heritage2026!Admin",
                "Rôle Administration, Accès C3, écriture autorisée.",
            ],
        ],
    )
    add_note(
        doc,
        "Mot de passe oublié :",
        "sur /sign-in, utiliser le lien proposé par Clerk (« Forgot password ») : un e-mail de réinitialisation est envoyé. L’application n’a pas de mot de passe maître.",
    )

    add_heading_custom(doc, "7. Ajouter un nouvel utilisateur", 1)
    add_para(
        doc,
        "Il n’existe pas encore d’écran « Créer un utilisateur » dans l’espace famille. Un utilisateur = un compte Clerk + une ligne User synchronisée à la première connexion.",
    )
    add_heading_custom(doc, "7.1 Depuis le site (cas habituel)", 2)
    add_steps(
        doc,
        [
            "Demander à la personne de s’inscrire elle-même (§ 5) avec son véritable e-mail familial.",
            "Vérifier qu’elle apparaît dans Prisma Studio, modèle User, après sa première visite de /espace.",
            "Attribuer le rôle et le statut généalogique (§ 4.1).",
            "Si la personne correspond à une fiche déjà présente (par ex. une personne de l’import), ne pas créer de doublon : le rattachement User.personId se fera lorsque l’écran de liaison sera disponible. En attendant, noter l’e-mail et l’identifiant de la fiche Personne (per-00xx ou person-…).",
        ],
    )
    add_heading_custom(doc, "7.2 Depuis le tableau de bord Clerk (administrateur technique)", 2)
    add_steps(
        doc,
        [
            "Ouvrir https://dashboard.clerk.com et sélectionner l’application liée au site.",
            "Aller dans Users puis Create user.",
            "Saisir l’e-mail et un mot de passe temporaire, puis le communiquer à la personne par un canal sûr.",
            "Lui demander de se connecter une fois sur /sign-in pour créer la ligne User locale.",
            "Promouvoir le rôle dans Prisma Studio comme au § 4.1.",
        ],
    )
    add_heading_custom(doc, "7.3 Exemple — ajouter un utilisateur comité", 2)
    add_steps(
        doc,
        [
            "Inscription : comite.exemple@toure-heritage.test / Heritage2026!Comite.",
            "Première connexion sur /espace : le rôle affiché est Invité.",
            "Dans Prisma Studio, User : role = COMMITTEE_MEMBER, genealogicalStatus = FAMILY_MEMBER.",
            "Reconnexion : le bandeau affiche « Comité restreint » et « Accès C3 ».",
            "La personne peut ouvrir /espace/personnes/nouvelle et créer une fiche C3 si nécessaire.",
        ],
    )

    add_heading_custom(doc, "Module utilisateur", 1)
    add_para(
        doc,
        "Les sections 8 à 15 s’adressent à la personne qui consulte et qui saisit : visiteur, membre de la Grande Famille ou descendant direct. Le comité et l’administration suivent aussi ces étapes, puis le module administrateur à partir du § 16.",
    )

    add_heading_custom(doc, "8. Insertion d’une nouvelle personne sur le site", 1)
    add_para(
        doc,
        "Condition : être connecté avec un rôle FAMILY_MEMBER, COMMITTEE_MEMBER ou ADMIN. Un Invité est renvoyé vers la liste s’il tente d’ouvrir /espace/personnes/nouvelle.",
    )
    add_heading_custom(doc, "8.1 Accéder au formulaire", 2)
    add_steps(
        doc,
        [
            "Se connecter et aller dans Espace famille → Personnes (/espace/personnes).",
            "Cliquer sur le bouton « Ajouter une personne » en haut à droite.",
            "La page « Ajouter une personne » s’ouvre. Le texte rappelle que la fiche naît en brouillon et qu’un homonyme déclenche un avertissement.",
        ],
    )
    add_heading_custom(doc, "8.2 Remplir la fiche", 2)
    add_table(
        doc,
        ["Champ à l’écran", "Obligation", "Exemple"],
        [
            ["Prénom", "Oui", "Karamoko"],
            ["Nom", "Oui (prérempli TOURÉ)", "TOURÉ"],
            ["Autres noms", "Non", "Karamoko Lanfia"],
            ["Genre", "Oui", "Masculin"],
            ["Naissance (texte)", "Non", "vers 1920"],
            ["Lieu de naissance", "Non", "Sanankoro"],
            ["Décès (texte)", "Non", "laisser vide si vivant"],
            ["Lieu de décès", "Non", ""],
            ["Occupation", "Non", "Cultivateur"],
            ["Ville de résidence", "Non", "Kankan"],
            ["Biographie", "Non", "Témoignage de… Source orale, à valider."],
            ["Branche", "Non", "Lignée de Lanfia"],
            ["Confidentialité", "Oui", "C1 · Grande Famille (défaut)"],
            ["Niveau de confiance", "Oui", "Moyenne"],
            ["Père", "Non", "Choisir une personne déjà créée, ou « Non documenté »"],
            ["Mère", "Non", "Idem"],
        ],
    )
    add_heading_custom(doc, "8.3 Enregistrer et traiter un homonyme", 2)
    add_steps(
        doc,
        [
            "Cliquer sur « Créer la fiche ».",
            "Si le prénom + nom existent déjà, une alerte « Homonymie possible » liste les fiches existantes avec un lien. Ne pas confirmer si c’est la même personne : ouvrir la fiche déjà là.",
            "Si c’est bien une autre personne (homonyme assumé), renvoyer le formulaire : la case cachée de confirmation est alors active et la nouvelle fiche est créée.",
            "L’application ouvre la fiche créée : /espace/personnes/[identifiant]. Statut de validation : Brouillon (DRAFT). Un journal d’audit PERSON_CREATED est écrit.",
            "Contrôler Parents, Conjoints, Enfants et Sources sur la fiche.",
        ],
    )
    add_note(
        doc,
        "Confidentialité :",
        "vous ne pouvez créer une fiche qu’à un niveau que vous avez le droit de lire. Un membre famille (C1) ne peut pas créer une fiche C3.",
    )

    add_heading_custom(doc, "9. Ajouter une personne et toute sa lignée", 1)
    add_para(
        doc,
        "L’arbre se construit du plus ancien vers le plus récent. On crée d’abord les parents, puis chaque enfant en les sélectionnant dans « Père » et « Mère ». On ne crée jamais un lien « parce que ça paraît logique ».",
    )
    add_heading_custom(doc, "9.1 Ordre obligatoire", 2)
    add_steps(
        doc,
        [
            "Repérer l’ancêtre le plus ancien déjà présent, ou le créer s’il n’existe pas. Exemple déjà en base : Lankafia TOURE → Samorigbé + Mâ Dianka → Lanfia TOURÉ + Mâ Sokona → Samory TOURÉ.",
            "Créer le conjoint s’il est nommé par une source, sans le fusionner avec un homonyme (Mâ Sokona ≠ Mâ Sokonassé ≠ Mâ Sokonossé TOURE).",
            "Pour chaque enfant de cette génération : Ajouter une personne, renseigner l’identité, choisir le père et la mère dans les listes, enregistrer.",
            "Passer à la génération suivante uniquement lorsque les parents existent déjà dans la liste.",
            "Vérifier l’arbre interne : /espace/arbre. Choisir le parent dans la liste, puis « Afficher ». L’arbre commence à cette personne. Un descendant n’apparaît que si la chaîne de pères et de mères est complète.",
            "« Lignée complète » montre l’ensemble depuis la racine. Les personnes sans lien vers cette racine restent dans « Non rattachés à la racine ». C’est normal tant que la filiation n’est pas documentée.",
        ],
    )
    add_heading_custom(doc, "9.2 Exemple guidé — ajouter un petit-fils de Lanfia", 2)
    add_para(doc, "Objectif pédagogique : enregistrer un fils de Samory, nommé ici à titre d’exemple « Massé Mamadi TOURE », sans prétendre valider historiquement ce rattachement si la source dit seulement « descendance de Kemo Lanfia ».")
    add_steps(
        doc,
        [
            "Se connecter en membre famille ou admin (ex. famille.exemple@toure-heritage.test).",
            "Aller dans Personnes et rechercher si « Massé Mamadi TOURE » existe déjà (fiche per-0024 importée de SRC-002, niveau C1). Si oui, ne pas recréer : ouvrir la fiche existante.",
            "Si la fiche existe déjà mais n’a pas Samory comme père : la modification du lien parent n’a pas encore d’écran dédié (voir § 10). Ne pas créer une seconde fiche.",
            "Si l’on crée un nouvel enfant encore absent : /espace/personnes/nouvelle.",
            "Prénom : Massé Mamadi. Nom : TOURE. Genre : Masculin. Branche : Lignée de Lanfia. Confidentialité : C1. Confiance : Faible ou Moyenne.",
            "Père : Samory TOURÉ (Almamy Samory TOURE). Mère : Non documenté (sauf source explicite).",
            "Biographie : indiquer la source (« Transcription SRC-002, à valider. Lien vers Samory non confirmé par REL-xxx. »).",
            "Créer la fiche. Ouvrir /espace/arbre, choisir Samory TOURÉ, puis « Afficher » : l’enfant apparaît dans sa descendance si le père choisi est bien Samory. Depuis la fiche, le bouton « Voir l’arbre depuis cette personne » ouvre la même vue.",
            "Répéter pour chaque frère ou sœur documenté, un par un. Ne jamais coller une liste entière sans contrôle d’homonymie.",
        ],
    )
    add_heading_custom(doc, "9.3 Exemple — nouvelle lignée collatérale (frère de Lanfia, si une source l’établit)", 2)
    add_steps(
        doc,
        [
            "Créer d’abord le père commun s’il n’existe pas (ici Samorigbé existe déjà).",
            "Ajouter la personne : Prénom du frère, Nom TOURÉ, Père = Samorigbé, Mère = Mâ Dianka si la source le dit.",
            "Cette personne n’apparaît pas sous Lanfia (elle n’est pas un enfant de Lanfia). Elle apparaît dans Personnes et sur la fiche de Samorigbé, rubrique Enfants.",
            "Pour afficher cette branche, ouvrir l’arbre et choisir Samorigbé ou Lankafia comme personne de départ. « Lignée complète » reste ancrée sur la racine officielle Lanfia (isRoot).",
        ],
    )
    add_note(
        doc,
        "Rappel qualité :",
        "les 22 autres noms de SRC-001 n’ont pas été rattachés automatiquement (relation « Non déterminée »). Les 139 noms SRC-002/003 sont liés à Lanfia seulement parce que le catalogue Excel écrit « Enfant de Kemo Lanfia (selon source) ». Toute correction de parent doit rester tracée.",
    )

    add_heading_custom(doc, "10. Modifier des informations déjà existantes", 1)
    add_heading_custom(doc, "10.1 Ce que le site permet aujourd’hui", 2)
    add_para(
        doc,
        "La fiche personne reste en lecture pour l’identité, la biographie et les liens déjà enregistrés : il n’y a pas de bouton « Modifier la fiche ». Depuis la fiche interne, un rôle d’écriture peut ajouter un conjoint (« Lier le conjoint » ou « Créer une nouvelle fiche et la lier »). Une source, un témoignage ou un document se dépose dans Héritage, pas en réécrivant la fiche. La correction d’un père, d’une mère ou d’un texte déjà publié reste dans Prisma Studio.",
    )
    add_heading_custom(doc, "10.2 Modifier son propre compte (identité de connexion)", 2)
    add_steps(
        doc,
        [
            "Une fois connecté, cliquer sur l’avatar Clerk en haut (UserButton).",
            "Ouvrir Manage account.",
            "Mettre à jour le prénom, le nom, l’e-mail ou le mot de passe.",
            "Enregistrer. L’application resynchronise e-mail, prénom et nom dans la table User à la prochaine requête authentifiée.",
        ],
    )
    add_heading_custom(doc, "10.3 Modifier une fiche généalogique existante (procédure actuelle)", 2)
    add_para(doc, "En attendant l’écran d’édition, la correction se fait dans Prisma Studio, par un administrateur :")
    add_steps(
        doc,
        [
            "À la racine du projet, lancer npm run db:studio.",
            "Ouvrir http://localhost:5555.",
            "Modèle Person : rechercher par firstName / lastName ou par id (person-lanfia-toure, per-0002, etc.).",
            "Corriger uniquement ce que la source autorise : orthographe, dates texte, lieu, biographie, confidentialityLevel, confidenceLevel, validationStatus.",
            "Ne pas fusionner deux lignes pour « faire plus propre » (Lanfia ≠ Lankafia).",
            "Pour changer un père ou une mère : modèle ParentChild. Créer une ligne (parentId, childId, parentRole FATHER ou MOTHER) ou supprimer un lien erroné. Unique : un seul père et une seule mère par type de rôle.",
            "Pour un conjoint : modèle SpouseRelationship. L’identifiant personIdA doit être inférieur (ordre alphabétique) à personIdB.",
            "Pour une nouvelle source, préférer l’écran Héritage (§ 14). Prisma Studio (Source puis PersonSource) ne sert qu’à corriger une source déjà enregistrée.",
            "Recharger la fiche sur le site pour contrôler le résultat.",
        ],
    )
    add_heading_custom(doc, "10.4 Exemple — corriger le lieu de naissance de Samory", 2)
    add_steps(
        doc,
        [
            "Prisma Studio → Person → id person-almamy-samory-toure.",
            "Champ birthPlace : déjà « Manyambaladugu ». Si une source plus précise arrive, remplacer le texte et noter la source dans PersonSource.",
            "Ne pas inventer une date exacte à la place de « vers 1830 » (birthDateText).",
            "Recharger http://localhost:3000/genealogie/person-almamy-samory-toure : le bandeau affiche « vers 1830 — 1900 ».",
        ],
    )
    add_heading_custom(doc, "10.5 Exemple — retirer un enfant rattaché par erreur à Lanfia", 2)
    add_steps(
        doc,
        [
            "Prisma Studio → ParentChild.",
            "Filtrer childId = identifiant de la personne (ex. per-0024).",
            "Supprimer uniquement le lien dont parentId = person-lanfia-toure si le comité décide que la source « Enfant de Kemo Lanfia » ne suffit pas.",
            "La fiche personne demeure ; elle passe dans les non rattachés de l’arbre. Aucune autre fiche n’est détruite (suppression en cascade seulement si l’on efface la Personne elle-même).",
        ],
    )

    add_heading_custom(doc, "11. Module utilisateur — consulter et rechercher", 1)
    add_para(doc, "Public : visiteur, invité, membre, descendant. Chacun ne voit que son niveau C0 à C3.")
    add_steps(
        doc,
        [
            "Liste interne : /espace/personnes — cartes filtrées selon le niveau du compte.",
            "Fiche interne : ouvrir une carte. Vérifier Parents, Conjoints, Enfants et Sources.",
            "Arbre : /espace/arbre ou /arbre. Choisir une personne, cliquer « Afficher ». L’affichage commence à cette personne. « Lignée complète » montre toute la lignée autorisée.",
            "Depuis une fiche, « Voir l’arbre depuis cette personne » ouvre directement cette descendance.",
            "Recherche interne : /espace/recherche — prénom, nom, autre nom ou ville. Les résultats ne dépassent jamais l’autorisation.",
            "Côté public, une personne n’apparaît que si son niveau est C0.",
        ],
    )

    add_heading_custom(doc, "12. Module utilisateur — lier un conjoint", 1)
    add_para(doc, "Condition : rôle FAMILY_MEMBER, COMMITTEE_MEMBER ou ADMIN. Le conjoint doit déjà exister, ou être créé puis lié. On ne fusionne pas les homonymes.")
    add_steps(
        doc,
        [
            "Ouvrir la fiche interne de la personne.",
            "Dans « Ajouter un conjoint », choisir une fiche existante, le type de lien (Mariage, Union, Autre), la date si elle est connue, et le lieu.",
            "Cliquer « Lier le conjoint ». Le lien apparaît dans Conjoints des deux fiches.",
            "Si la conjointe ou le conjoint n’existe pas encore : « Créer une nouvelle fiche et la lier ». Le formulaire de création s’ouvre avec le rattachement prévu.",
        ],
    )

    add_heading_custom(doc, "13. Module utilisateur — vie de famille", 1)
    add_para(doc, "Réunions, événements, discussions et notifications exigent un accès C1 (membre, descendant direct avec ce niveau, comité, administration). Un invité C0 voit le refus d’accès.")
    add_heading_custom(doc, "13.1 Réunions et événements", 2)
    add_steps(
        doc,
        [
            "Menu Réunions. Renseigner Titre, Date et heure, Lieu, Ordre du jour, puis « Publier la réunion ».",
            "Menu Événements. Choisir le type (mariage, funérailles, patrimoine, autre), puis « Publier l’événement ».",
        ],
    )
    add_heading_custom(doc, "13.2 Discussions et vote", 2)
    add_steps(
        doc,
        [
            "Menu Discussions. Saisir le sujet et le texte, puis « Ouvrir la discussion ».",
            "Ouvrir le sujet. Commenter avec « Commenter ».",
            "Voter une seule fois, de façon nominative : Pour, Contre ou Abstention. Le décompte s’affiche sur la discussion.",
        ],
    )

    add_heading_custom(doc, "14. Module utilisateur — déposer une pièce d’héritage", 1)
    add_para(
        doc,
        "Un témoignage oral n’est pas un fait établi. Le dépôt crée un brouillon. Seul le comité peut l’approuver, et seulement avec un procès-verbal. La page Histoire n’affiche ensuite que les pièces de niveau C0 déjà publiées.",
    )
    add_steps(
        doc,
        [
            "Menu Héritage. Onglets : Sources, Témoignages, Documents, Validation.",
            "Source : titre, type, auteur, adresse, personne liée, description, confidentialité et confiance, puis « Enregistrer le brouillon ».",
            "Témoignage : titre, personne qui témoigne, lieu, période, étape du recueil (brut, transcrit ou relu), résumé, transcription, audio ou vidéo. La confiance de départ est faible.",
            "Document : nature, date indiquée, auteur, lieu de conservation, fichier, personne liée.",
            "Choisir C0 seulement si la pièce pourra être publique après validation. Le défaut est C1, Grande Famille.",
            "Sur la carte du brouillon, ou après un rejet, cliquer « Soumettre au comité ».",
        ],
    )

    add_heading_custom(doc, "15. Contrôles avant de quitter une saisie", 1)
    add_table(
        doc,
        ["Contrôle", "Question à se poser"],
        [
            ["Source", "Ai-je indiqué d’où vient l’information (SRC-001, témoignage, archive) ?"],
            ["Témoignage", "Ai-je laissé le récit en brouillon, sans le présenter comme un fait ?"],
            ["Homonyme", "Ai-je ouvert les fiches existantes avant d’en créer une autre ?"],
            ["Lien", "Le père, la mère ou le conjoint est-il écrit dans une source, ou l’ai-je déduit ?"],
            ["Fusion", "Ai-je évité de mélanger Lanfia et Lankafia, Sokona et Sokonassé ?"],
            ["Confidentialité", "Le niveau C0 à C3 est-il justifié ? Une liste familiale brute n’est pas publique."],
            ["Validation", "La fiche ou la pièce est-elle restée en brouillon tant que le comité n’a pas statué ?"],
            ["Audit", "La création a-t-elle eu lieu sous un compte nommé, sans partage de mot de passe ?"],
        ],
    )

    add_heading_custom(doc, "16. Module administrateur — promouvoir un profil", 1)
    add_para(
        doc,
        "Il n’existe pas d’écran « Créer un utilisateur » ni « Changer le rôle » dans l’espace famille. L’administration technique passe par Clerk pour le compte, puis par Prisma Studio pour le rôle. La procédure détaillée est au § 4.1 et au § 7.",
    )
    add_steps(
        doc,
        [
            "Faire inscrire la personne (§ 5) ou la créer dans le tableau de bord Clerk (§ 7.2).",
            "Après sa première visite de /espace, ouvrir Prisma Studio (npm run db:studio, http://localhost:5555), modèle User.",
            "Attribuer role et genealogicalStatus. Exemple comité : COMMITTEE_MEMBER et FAMILY_MEMBER.",
            "Demander une reconnexion. Le tableau de bord doit afficher le nouveau rôle et l’accès C1, C2 ou C3.",
            "Le premier administrateur peut être promu par BOOTSTRAP_ADMIN_CLERK_USER_ID tant que son rôle est encore GUEST.",
            "Ne pas lier User.personId à la main sans décision : l’écran de rattachement compte–fiche n’existe pas encore. Noter l’e-mail et l’identifiant de la fiche.",
        ],
    )

    add_heading_custom(doc, "17. Module administrateur — valider, procès-verbal et publication", 1)
    add_para(doc, "Réservé au comité restreint et à l’administration. Un membre peut soumettre et lire le suivi des pièces qu’il a le droit de voir. Il ne tranche pas.")
    add_heading_custom(doc, "17.1 Rédiger le procès-verbal", 2)
    add_steps(
        doc,
        [
            "Ouvrir Héritage → Validation.",
            "Dans « Rédiger un procès-verbal », saisir le titre, la date, les membres présents, la décision et la mention de signature.",
            "Joindre l’adresse du fichier si le PV numérisé existe.",
            "Cocher « Publier ce procès-verbal » seulement s’il peut être public : il passe alors en C0. Sinon il reste interne (C1).",
            "Cliquer « Enregistrer le procès-verbal ».",
        ],
    )
    add_heading_custom(doc, "17.2 Trancher une demande", 2)
    add_steps(
        doc,
        [
            "Sur la demande en attente, choisir la décision : Demander une révision, Rejeter, ou Approuver.",
            "L’approbation exige un procès-verbal dans la liste. Sans PV, l’enregistrement est refusé.",
            "Écrire le motif, puis « Enregistrer la décision ».",
            "La révision renvoie la pièce en brouillon. Le rejet la marque rejetée. L’auteur peut la soumettre de nouveau après correction.",
            "« Commenter » place une demande soumise en examen et conserve le commentaire du comité.",
        ],
    )
    add_heading_custom(doc, "17.3 Publier sur la page Histoire", 2)
    add_steps(
        doc,
        [
            "La publication n’est proposée que si la pièce est déjà approuvée et de niveau C0.",
            "Cliquer « Publier sur la page Histoire ».",
            "Ouvrir /histoire sans être connecté : la source, le témoignage, le document ou le procès-verbal public doit apparaître. Un témoignage publié reste présenté comme un témoignage.",
            "Une pièce C1, C2 ou C3 approuvée reste dans l’espace famille. Elle ne va pas sur Histoire.",
        ],
    )
    add_note(
        doc,
        "Exemple :",
        "un membre dépose un témoignage en C0, confiance faible, puis « Soumettre au comité ». Le comité rédige le PV, choisit Approuver et ce PV, puis publie. Le visiteur voit le récit sur Histoire, avec la mention qu’il reste un témoignage.",
    )

    add_heading_custom(doc, "18. Module administrateur — caisse", 1)
    add_para(doc, "La caisse est réservée au comité et à l’administration. Un membre de la Grande Famille voit « Caisse familiale » refusée.")
    add_steps(
        doc,
        [
            "Ouvrir /espace/caisse.",
            "Cotisation : contributeur, montant en GNF, date, objet, puis « Ajouter ».",
            "Dépense : libellé, montant, date, catégorie, puis « Ajouter ».",
            "Contrôler le solde affiché. Chaque écriture est auditable. Ne pas saisir un montant sans pièce ou décision du comité.",
        ],
    )

    add_heading_custom(doc, "19. Ce que l’écran ne fait pas encore", 1)
    add_para(
        doc,
        "Le site permet l’inscription, la connexion, la consultation filtrée, l’arbre à partir d’une personne choisie, la création d’une fiche avec père et mère, le conjoint, la vie de famille, le vote nominatif, le dépôt d’héritage, la validation par le comité et la publication C0. Il ne permet pas encore, à l’écran, de modifier une fiche déjà créée, de changer un rôle, ni de lier un compte à une fiche Personne. Ces trois opérations restent celles des § 7, 10 et 16.",
    )
    add_para(
        doc,
        "Fin des modules. Toute évolution d’écran (bouton Modifier, administration des rôles) devra mettre à jour ce guide en même temps que l’application.",
        italic=True,
        color=MUTED,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
