"""Génère le guide de publication Hostinger au format Word."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "Procedure_publication_Hostinger_TOURE_FAMILY_HERITAGE.docx"

UMBER = RGBColor(0x5C, 0x3A, 0x21)
MUTED = RGBColor(0x5A, 0x4A, 0x3A)
BLACK = RGBColor(0x1F, 0x1A, 0x14)


def set_run(run, *, size=11, bold=False, color=BLACK, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, italic=False, color=BLACK, space_after=8, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading_custom(doc, text, level):
    sizes = {1: 20, 2: 15, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run(run, size=sizes[level], bold=True, color=UMBER)
    return p


def add_steps(doc, steps):
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.space_after = Pt(6)
        num = p.add_run(f"{i}. ")
        set_run(num, size=11, bold=True, color=UMBER)
        body = p.add_run(step)
        set_run(body, size=11)


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {qn("w:fill"): "5C3A21", qn("w:val"): "clear"})
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_header(cell)
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(header)
        set_run(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for r_idx, row in enumerate(rows, 1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(value)
            set_run(run, size=10)
    doc.add_paragraph()
    return table


def add_note(doc, title, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.3)
    t = p.add_run(f"{title} ")
    set_run(t, size=11, bold=True, color=UMBER)
    b = p.add_run(text)
    set_run(b, size=11, italic=True, color=MUTED)


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("TOURÉ FAMILY HERITAGE — Publication Hostinger — Document interne")
    set_run(fr, size=8, color=MUTED)

    add_para(doc, "TOURÉ FAMILY HERITAGE", size=12, bold=True, color=UMBER, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_para(
        doc,
        "Procédure de publication du site sur Hostinger",
        size=22,
        bold=True,
        color=UMBER,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=8,
    )
    add_para(
        doc,
        "Guide étape par étape pour mettre en ligne l’application Next.js (généalogie, Clerk, base Neon) depuis hPanel.",
        size=12,
        italic=True,
        color=MUTED,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=16,
    )
    add_para(
        doc,
        f"Version 1.0  ·  {date.today().strftime('%d/%m/%Y')}  ·  Dépôt GitHub https://github.com/hassane7toure-hub/genealogievf",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )

    add_heading_custom(doc, "1. Objet du document", 1)
    add_para(
        doc,
        "Ce document décrit, dans l’ordre, comment publier TOURÉ FAMILY HERITAGE sur Hostinger pour qu’il soit accessible sur Internet avec un nom de domaine. Il correspond à l’application actuelle : Next.js 16, authentification Clerk, base PostgreSQL Neon, arbre généalogique et espace famille.",
    )
    add_note(
        doc,
        "Important :",
        "GitHub stocke le code, il ne publie pas le site. Un hébergement Hostinger classique (PHP, dossier public_html, fichiers HTML) ne peut pas faire tourner cette application. Il faut un hébergement Node.js (application web) ou un VPS.",
    )

    add_heading_custom(doc, "2. Ce qu’il ne faut pas faire", 1)
    add_steps(
        doc,
        [
            "Ne pas glisser le dossier du projet dans le Gestionnaire de fichiers Hostinger : Next.js n’est pas un site HTML statique.",
            "Ne pas utiliser le Git Hostinger « simple copie de fichiers » : il ne lance pas npm run build ni le serveur Node.",
            "Ne pas exporter le site en pages statiques : Clerk, Prisma et l’espace famille exigent un serveur.",
            "Ne pas commiter le fichier .env.local : les mots de passe et clés secrètes restent dans le tableau de bord Hostinger.",
        ],
    )

    add_heading_custom(doc, "3. Offre Hostinger à choisir", 1)
    add_para(
        doc,
        "Deux chemins existent. Le chemin recommandé pour ce projet est l’application Node.js dans hPanel, sans ligne de commande.",
    )
    add_table(
        doc,
        ["Chemin", "Offre Hostinger", "Quand l’utiliser"],
        [
            [
                "Recommandé — Application Node.js",
                "Business Web Hosting, ou Cloud Startup / Professional / Enterprise",
                "Publication depuis GitHub, build automatique, redémarrage géré par Hostinger",
            ],
            [
                "Avancé — VPS",
                "VPS KVM Ubuntu",
                "Vous préférez SSH, Nginx et PM2. Plus de contrôle, plus technique",
            ],
            [
                "Inadapté",
                "Hébergement web Premium / Single (PHP)",
                "Suffit pour WordPress, pas pour Next.js",
            ],
        ],
    )
    add_note(
        doc,
        "Vérification dans hPanel :",
        "après connexion, allez dans Sites web → Ajouter un site. Si vous voyez « Application web Node.js », votre offre est compatible. Sinon, passez à l’offre Business/Cloud ou utilisez un VPS (annexe).",
    )

    add_heading_custom(doc, "4. Préparer le code sur GitHub", 1)
    add_para(
        doc,
        "Hostinger lit le code depuis GitHub. Le dépôt du projet est déjà : https://github.com/hassane7toure-hub/genealogievf",
    )
    add_steps(
        doc,
        [
            "Ouvrez le projet sur votre ordinateur (Cursor ou l’explorateur de fichiers).",
            "Vérifiez que le site fonctionne en local : npm run dev puis http://localhost:3000.",
            "Vérifiez que le build passe : npm run build. S’il échoue, ne déployez pas.",
            "Ouvrez GitHub Desktop ou le terminal, puis poussez la branche master (ou main) vers GitHub.",
            "Sur github.com/hassane7toure-hub/genealogievf, confirmez que les derniers fichiers (espace famille, conjoints, etc.) sont bien visibles.",
            "Ne poussez jamais .env, .env.local ni les mots de passe.",
        ],
    )
    add_note(
        doc,
        "Branche :",
        "Hostinger déploie la branche par défaut du dépôt (souvent master). Tous les commits que vous voulez en ligne doivent être sur cette branche.",
    )

    add_heading_custom(doc, "5. Préparer la base Neon (déjà utilisée en local)", 1)
    add_para(
        doc,
        "La généalogie et l’espace famille sont dans PostgreSQL Neon. Hostinger n’héberge pas cette base : le site distant se connecte à la même base (ou à une copie production).",
    )
    add_steps(
        doc,
        [
            "Ouvrez https://console.neon.tech et connectez-vous.",
            "Ouvrez le projet de la Grande Famille TOURÉ.",
            "Copiez la chaîne de connexion groupée (pooled), qui commence par postgresql:// et contient sslmode=require.",
            "Gardez-la pour l’étape « Variables d’environnement ». Ne la collez pas dans GitHub.",
            "Si vous voulez séparer local et production, créez une deuxième base Neon, puis relancez npx prisma db push et npm run db:seed uniquement sur cette base.",
        ],
    )
    add_note(
        doc,
        "Conseil :",
        "pour un premier lancement, réutiliser la base actuelle est plus simple. Passez à une base production dédiée quand le site est stable.",
    )

    add_heading_custom(doc, "6. Préparer Clerk pour le domaine public", 1)
    add_para(
        doc,
        "En local, Clerk utilise des clés de développement (pk_test_ / sk_test_). Sur Internet, il faut des clés de production et autoriser votre nom de domaine.",
    )
    add_steps(
        doc,
        [
            "Ouvrez https://dashboard.clerk.com et sélectionnez l’application TOURÉ FAMILY HERITAGE.",
            "Passez sur l’instance Production (pas Development).",
            "Dans API Keys, copiez la Publishable Key (pk_live_…) et la Secret Key (sk_live_…).",
            "Dans Configure → Domains (ou Paths), ajoutez le domaine Hostinger, par exemple www.votrefamille.com et votrefamille.com.",
            "Laissez les chemins /sign-in et /sign-up : ils correspondent déjà à l’application.",
            "Les adresses de redirection après connexion restent /espace.",
            "Si vous testez d’abord l’URL temporaire Hostinger, ajoutez aussi cette URL dans les domaines autorisés Clerk.",
        ],
    )
    add_note(
        doc,
        "Sans cette étape :",
        "le site peut s’afficher mais la connexion échoue (Clerk refuse le domaine, ou les clés test bloquent la production).",
    )

    add_heading_custom(doc, "7. Créer le site Node.js dans hPanel", 1)
    add_para(
        doc,
        "Après la souscription, Hostinger ouvre souvent l’écran « Choisissez une méthode de migration pour votre site ». Cet écran sert à importer WordPress, cPanel ou des fichiers HTML/PHP. Ce n’est pas le chemin de TOURÉ FAMILY HERITAGE.",
    )
    add_steps(
        doc,
        [
            "Ne laissez pas « Utiliser les identifiants » (WordPress / cPanel) : ce n’est pas un site WordPress.",
            "Ne choisissez pas « Importer les fichiers de sauvegarde » (HTML, CSS, PHP) : Next.js n’est pas un site statique.",
            "Ne cliquez pas sur Suivant avec l’une de ces deux cartes.",
            "Cliquez sur le lien vert en bas de l’écran : Utilisateurs avancés — Déployez votre Application web Node.js.",
            "À l’écran suivant, choisissez Importer un dépôt Git (Import Git repository).",
            "Cliquez sur Connecter avec GitHub. Une fenêtre demande d’installer l’application GitHub Hostinger.",
            "Autorisez l’accès au compte hassane7toure-hub et au dépôt genealogievf.",
            "Sélectionnez genealogievf puis validez.",
            "Si le dépôt n’apparaît pas, cliquez sur Actualiser les dépôts (Refresh repositories).",
        ],
    )
    add_note(
        doc,
        "Si vous avez déjà cliqué Suivant par erreur :",
        "revenez avec le bouton Retour, puis cliquez sur Application web Node.js. Si le site PHP a déjà été créé, ouvrez Sites web → Ajouter un site et choisissez directement Node.js.",
    )

    add_heading_custom(doc, "8. Réglages de compilation à vérifier", 1)
    add_para(
        doc,
        "Hostinger détecte souvent Next.js tout seul. Contrôlez chaque champ avant de lancer le premier déploiement.",
    )
    add_table(
        doc,
        ["Champ hPanel", "Valeur à indiquer"],
        [
            ["Préréglage de framework", "Next.js"],
            ["Branche", "master (ou la branche par défaut du dépôt)"],
            ["Version Node.js", "22 (ou 20 LTS si 22 n’est pas proposé)"],
            ["Gestionnaire de paquets", "npm"],
            ["Commande d’installation", "laisser le défaut (npm ci ou npm install)"],
            ["Commande de build", "npm run build"],
            ["Dossier de sortie", ".next"],
            ["Répertoire racine", "vide (le package.json est à la racine du dépôt)"],
            ["Commande de démarrage", "npm run start -- -p $PORT"],
            ["Fichier d’entrée", "laisser vide pour Next.js, sauf si Hostinger l’exige"],
        ],
    )
    add_note(
        doc,
        "Démarrage :",
        "Hostinger fournit le port via $PORT. La commande officielle pour Next.js est npm run start -- -p $PORT. Sans le port, le site peut compiler mais rester inaccessible.",
    )

    add_heading_custom(doc, "9. Renseigner les variables d’environnement", 1)
    add_para(
        doc,
        "Dans l’écran de déploiement, ouvrez Variables d’environnement (ou, après coup, le menu du site → Environment variables). Ajoutez-les une par une, ou importez un fichier .env sans jamais le commiter.",
    )
    add_table(
        doc,
        ["Nom de la variable", "Exemple / source", "Obligatoire"],
        [
            ["DATABASE_URL", "Chaîne postgresql://… copiée depuis Neon", "Oui"],
            ["NEXT_PUBLIC_CLERK_PUBLISHABLE_KEY", "pk_live_… depuis Clerk Production", "Oui"],
            ["CLERK_SECRET_KEY", "sk_live_… depuis Clerk Production", "Oui"],
            ["NEXT_PUBLIC_CLERK_SIGN_IN_URL", "/sign-in", "Oui"],
            ["NEXT_PUBLIC_CLERK_SIGN_UP_URL", "/sign-up", "Oui"],
            ["NEXT_PUBLIC_CLERK_SIGN_IN_FALLBACK_REDIRECT_URL", "/espace", "Oui"],
            ["NEXT_PUBLIC_CLERK_SIGN_UP_FALLBACK_REDIRECT_URL", "/espace", "Oui"],
            ["BOOTSTRAP_ADMIN_CLERK_USER_ID", "Identifiant Clerk de l’administrateur (optionnel)", "Non"],
        ],
    )
    add_steps(
        doc,
        [
            "Cliquez sur Ajouter une variable d’environnement pour chaque ligne du tableau.",
            "Collez les vraies valeurs (pas pk_test_ en production).",
            "Enregistrez. Hostinger relance un déploiement pour que les valeurs soient prises en compte à la compilation et à l’exécution.",
            "Les clés NEXT_PUBLIC_… doivent être présentes AVANT le build : elles sont inscrites dans le site généré.",
        ],
    )
    add_note(
        doc,
        "Assistant base Hostinger :",
        "le bouton « Connecter une base » propose Supabase ou MongoDB. Ne l’utilisez pas : ce projet utilise déjà Neon. Collez seulement DATABASE_URL.",
    )

    add_heading_custom(doc, "10. Lancer le déploiement", 1)
    add_steps(
        doc,
        [
            "Vérifiez framework, branche, Node 22, commande npm run build et les variables.",
            "Cliquez sur Déployer (Deploy).",
            "Restez sur la page : les journaux de compilation s’affichent en direct.",
            "Attendez le statut réussi (build vert). La première fois peut durer plusieurs minutes (npm install, prisma generate, next build).",
            "Si le build échoue, ouvrez Déploiements, lisez le journal, corrigez, puis redéployez. Hostinger propose aussi une analyse automatique de l’échec.",
        ],
    )
    add_para(
        doc,
        "Les fichiers compilés sont placés par Hostinger sous ~/domains/{votre-domaine}/hbuilds/current/nodejs. Vous n’avez pas à les copier à la main. Un fichier .htaccess est généré dans public_html pour diriger les visiteurs vers Node.js : ne le modifiez pas.",
    )

    add_heading_custom(doc, "11. Relier le nom de domaine et le HTTPS", 1)
    add_steps(
        doc,
        [
            "Dans hPanel, ouvrez Domaines. Si le domaine est acheté chez Hostinger, sélectionnez-le pour ce site Node.js.",
            "Si le domaine est ailleurs (registrar externe), créez un enregistrement A vers l’adresse IP indiquée par Hostinger, et un CNAME www vers le domaine principal, selon les consignes hPanel.",
            "Attendez la propagation DNS (parfois quelques minutes, parfois quelques heures).",
            "Hostinger active en général le certificat SSL (cadenas HTTPS) automatiquement. Sinon, ouvrez SSL et demandez un certificat Let’s Encrypt.",
            "Testez https://www.votredomaine.com et non seulement http://.",
            "Retournez dans Clerk et confirmez que ce domaine exact (avec et sans www si les deux existent) est autorisé.",
        ],
    )

    add_heading_custom(doc, "12. Vérifier le site en ligne", 1)
    add_para(doc, "Ouvrez le site comme un visiteur, puis comme administrateur.")
    add_steps(
        doc,
        [
            "Accueil : le titre TOURÉ FAMILY HERITAGE et le portrait s’affichent.",
            "Histoire, Généalogie, Arbre, Recherche : pages publiques C0, sans compte.",
            "Bouton Rejoindre / Connexion : le formulaire Clerk s’ouvre (plus d’erreur « Failed to load Clerk JS »).",
            "Après connexion : /espace affiche le tableau de bord, pas une page 404.",
            "Arbre interne : les descendants de Lanfia et de Samory apparaissent.",
            "Fiche d’une personne : les conjoints déjà enregistrés sont listés ; le formulaire « Ajouter un conjoint » est visible pour un compte autorisé à écrire.",
            "Si une page reste blanche, ouvrez le site → Journaux d’exécution (Runtime Logs) dans hPanel.",
        ],
    )
    add_table(
        doc,
        ["Page", "Adresse une fois le domaine posé"],
        [
            ["Accueil", "https://votredomaine.com/"],
            ["Histoire", "https://votredomaine.com/histoire"],
            ["Arbre public", "https://votredomaine.com/arbre"],
            ["Inscription", "https://votredomaine.com/sign-up"],
            ["Espace famille", "https://votredomaine.com/espace"],
        ],
    )

    add_heading_custom(doc, "13. Publier une mise à jour plus tard", 1)
    add_steps(
        doc,
        [
            "Modifiez le site en local, testez avec npm run dev.",
            "Lancez npm run build pour confirmer qu’il n’y a pas d’erreur.",
            "Poussez vers GitHub (branche connectée à Hostinger).",
            "Hostinger reçoit le webhook GitHub, installe, compile et redémarre tout seul.",
            "Suivez l’avancement dans Sites web → votre site → Déploiements.",
            "Si les variables d’environnement changent, enregistrez-les dans hPanel : cela relance un déploiement.",
            "Pour un simple redémarrage sans nouveau code, cliquez sur le badge En cours (Running) puis Redémarrer.",
        ],
    )

    add_heading_custom(doc, "14. Problèmes fréquents", 1)
    add_table(
        doc,
        ["Symptôme", "Cause probable", "Que faire"],
        [
            [
                "Build échoue sur prisma generate",
                "Node trop ancien ou postinstall bloqué",
                "Forcer Node 22 et la commande npm run build",
            ],
            [
                "Site compilé mais page d’erreur au démarrage",
                "DATABASE_URL ou clés Clerk absentes",
                "Runtime Logs + vérifier les variables, puis enregistrer (redéploie)",
            ],
            [
                "Connexion Clerk impossible",
                "Domaine absent dans Clerk, ou clés test en production",
                "Ajouter le domaine, passer en pk_live_ / sk_live_",
            ],
            [
                "403 après un redéploiement",
                ".htaccess modifié à la main",
                "Redéployer pour le régénérer, ne plus l’éditer",
            ],
            [
                "Arbre vide / base indisponible",
                "Mauvaise DATABASE_URL ou base injoignable",
                "Recopier l’URL Neon pooled, sslmode=require",
            ],
            [
                "Le dépôt n’apparaît pas",
                "Application GitHub Hostinger mal autorisée",
                "Reconnecter GitHub et Actualiser les dépôts",
            ],
        ],
    )

    add_heading_custom(doc, "15. Annexe — publication sur VPS Hostinger", 1)
    add_para(
        doc,
        "Utilisez cette annexe seulement si votre offre n’a pas « Application web Node.js ». Il faut un VPS Ubuntu 24.04, un accès SSH, et plus de manipulations.",
    )
    add_steps(
        doc,
        [
            "Dans hPanel, ouvrez le VPS, notez l’adresse IP et le mot de passe root.",
            "Connectez-vous en SSH : ssh root@ADRESSE_IP",
            "Installez Node.js 22 depuis NodeSource (pas le paquet Ubuntu par défaut), puis git, nginx et pm2.",
            "Clonez le dépôt dans /var/www/genealogievf.",
            "Créez le fichier .env.local (ou .env) sur le serveur avec les mêmes variables que le tableau de l’étape 9.",
            "Exécutez : npm ci puis npm run build.",
            "Démarrez : pm2 start npm --name genealogievf -- start",
            "Enregistrez le redémarrage automatique : pm2 startup puis pm2 save.",
            "Configurez Nginx en proxy vers http://127.0.0.1:3000 pour votre nom de domaine.",
            "Activez HTTPS avec certbot --nginx.",
            "Pour une mise à jour : git pull, npm ci, npm run build, pm2 restart genealogievf.",
        ],
    )
    add_note(
        doc,
        "Sécurité VPS :",
        "changez le mot de passe root, créez un utilisateur non-root, et n’ouvrez que les ports 22, 80 et 443.",
    )

    add_heading_custom(doc, "16. Liste de contrôle avant d’annoncer le site", 1)
    add_steps(
        doc,
        [
            "Offre Hostinger avec Application web Node.js (ou VPS opérationnel).",
            "Code à jour poussé sur GitHub.",
            "npm run build réussi en local.",
            "DATABASE_URL Neon renseignée dans hPanel.",
            "Clés Clerk de production + domaine autorisé.",
            "Déploiement Hostinger vert.",
            "HTTPS actif sur le nom de domaine.",
            "Accueil, arbre public et connexion testés.",
            "Un compte administrateur peut ouvrir /espace.",
            "Aucune clé secrète n’est visible dans le dépôt GitHub.",
        ],
    )

    add_para(
        doc,
        "Document interne — TOURÉ FAMILY HERITAGE. La publication n’autorise pas à inventer des liens généalogiques ni à fusionner des homonymes.",
        italic=True,
        color=MUTED,
        space_after=0,
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
